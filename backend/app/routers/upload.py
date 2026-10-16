"""Dataset upload router: a CSV/Excel file -> Easy Builder draft payload."""
from typing import Any, Dict, List
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile

from app.auth import get_current_user

router = APIRouter(tags=["upload"], dependencies=[Depends(get_current_user)])

GUIDE_COLUMNS = ["teacher_name", "course_name", "credit_hours", "student_group", "preferred_time"]
_BANDS = ("Morning", "Afternoon", "Evening")


def _rows_from_workbook(data: bytes) -> List[List[str]]:
    from openpyxl import load_workbook
    import io

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([(str(c).strip() if c is not None else "") for c in row])
    return [r for r in rows if any(r)]


def _validate_schedule_rows(rows: List[List[str]]) -> tuple:
    """Validate guide-format rows -> (courses, errors). Row numbers are 1-based incl. header."""
    courses: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    seen = set()
    start = 0
    if rows and [c.lower() for c in rows[0][:2]] == ["teacher_name", "course_name"]:
        start = 1  # header row
    for i in range(start, len(rows)):
        r = (rows[i] + ["", "", "", "", ""])[:5]
        teacher, course, credits, group, band = [c.strip() for c in r]
        row_no = i + 1
        problems = []
        if not teacher:
            problems.append("Teacher name is missing")
        if not course:
            problems.append("Course name is missing")
        try:
            ch = int(float(credits)) if credits else 3
            if ch < 1 or ch > 4:
                problems.append("Credit hours must be 1 to 4")
        except (ValueError, TypeError):
            problems.append("Credit hours must be 1 to 4")
            ch = 3
        band_norm = band.capitalize()
        if band and band_norm not in _BANDS:
            problems.append("Preferred time must be Morning, Afternoon, Evening, or blank")
            band_norm = ""
        if teacher and course and (teacher.lower(), course.lower()) in seen:
            problems.append(f"Duplicate: teacher '{teacher}' already teaches '{course}'")
        if problems:
            errors.append({"row": row_no, "message": "; ".join(problems)})
            continue
        seen.add((teacher.lower(), course.lower()))
        courses.append({
            "name": course,
            "teacher": teacher,
            "credit_hours": ch,
            "capacity_required": 30,
            "student_group": group,
            "preferred_time": band_norm,
        })
    return courses, errors


@router.post("/upload/file")
async def upload_file(file: UploadFile = File(...)):
    """Upload a CSV or Excel dataset; returns preview, row errors, and draft payload."""
    filename = (file.filename or "").lower()
    if not filename.endswith((".csv", ".xlsx", ".xls")):
        raise HTTPException(
            status_code=400,
            detail="Upload a .csv or .xlsx file with columns: " + ", ".join(GUIDE_COLUMNS),
        )
    try:
        raw = await file.read()
        rows = _rows_from_workbook(raw) if filename.endswith(("xlsx", "xls")) else [
            [c.strip() for c in line.split(",")]
            for line in raw.decode("utf-8-sig").splitlines()
            if line.strip() and not line.strip().startswith("#")
        ]
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read file: {exc}")
    courses, errors = _validate_schedule_rows(rows)
    teachers = []
    for c in courses:
        if c["teacher"] not in [t["name"] for t in teachers]:
            teachers.append({"name": c["teacher"], "max_hours_per_day": 6,
                             "unavailable_days": []})
    return {
        "filename": file.filename,
        "columns": GUIDE_COLUMNS,
        "total_rows": max(0, len(rows) - (1 if rows and [c.lower() for c in rows[0][:2]] == ["teacher_name", "course_name"] else 0)),
        "preview": courses[:10],
        "errors": errors,
        "teachers": teachers,
        "courses": courses,
    }
